"""Сервис экспорта: генерация DOCX / PDF / TXT из результатов оптимизации."""

from __future__ import annotations

import io
import logging
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

logger = logging.getLogger(__name__)


def generate_docx(rewritten_data: dict[str, Any] | None, *, raw_text: str | None = None) -> bytes:
    """Генерация DOCX-файла из результатов оптимизации.

    Args:
        rewritten_data: Структурированные данные от LLM (JSON).
        raw_text: Fallback — текстовый ответ LLM, если JSON невалиден.

    Returns:
        DOCX-файл в виде bytes.
    """
    doc = Document()

    # Стили
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    if rewritten_data:
        _build_structured_docx(doc, rewritten_data)
    elif raw_text:
        _build_plain_docx(doc, raw_text)
    else:
        doc.add_paragraph('Данные для экспорта отсутствуют.')

    # Сохранение в bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _build_structured_docx(doc: Document, data: dict[str, Any]) -> None:
    """Генерация структурированного DOCX из JSON-данных LLM."""
    # Профессиональное саммари
    summary = data.get('summary', '')
    if summary:
        heading = doc.add_heading('Профессиональное резюме', level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph(summary)
        p.paragraph_format.space_after = Pt(12)

    # Опыт работы
    experience = data.get('experience', [])
    if experience:
        doc.add_heading('Опыт работы', level=2)
        for exp in experience:
            position = exp.get('position', '')
            company = exp.get('company', '')
            period = exp.get('period', '')

            p = doc.add_paragraph()
            run = p.add_run(position)
            run.bold = True
            run.font.size = Pt(12)

            if company or period:
                p2 = doc.add_paragraph()
                if company:
                    p2.add_run(company).italic = True
                if company and period:
                    p2.add_run(f' | {period}')
                elif period:
                    p2.add_run(period)

            achievements = exp.get('achievements', [])
            for ach in achievements:
                doc.add_paragraph(ach, style='List Bullet')

            doc.add_paragraph()  # Отступ

    # Образование
    education = data.get('education', [])
    if education:
        doc.add_heading('Образование', level=2)
        for edu in education:
            institution = edu.get('institution', '')
            degree = edu.get('degree', '')
            specialization = edu.get('specialization', '')
            year = edu.get('year', '')

            p = doc.add_paragraph()
            run = p.add_run(institution)
            run.bold = True

            details = []
            if degree:
                details.append(degree)
            if specialization:
                details.append(specialization)
            if year:
                details.append(str(year))
            if details:
                doc.add_paragraph(', '.join(details))

    # Навыки
    skills = data.get('skills', [])
    if skills:
        doc.add_heading('Навыки', level=2)
        doc.add_paragraph(', '.join(skills))


def _build_plain_docx(doc: Document, text: str) -> None:
    """Генерация простого DOCX из текста."""
    heading = doc.add_heading('Оптимизированное резюме', level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for paragraph in text.split('\n'):
        stripped = paragraph.strip()
        if stripped:
            doc.add_paragraph(stripped)


# ---------------------------------------------------------------------------
# PDF generation (reportlab)
# ---------------------------------------------------------------------------

def generate_pdf(rewritten_data: dict[str, Any] | None, *, raw_text: str | None = None) -> bytes:
    """Генерация PDF-файла из результатов оптимизации.

    Args:
        rewritten_data: Структурированные данные от LLM (JSON).
        raw_text: Fallback — текстовый ответ LLM, если JSON невалиден.

    Returns:
        PDF-файл в виде bytes.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()
    # Custom styles
    styles.add(ParagraphStyle(
        'ResumeTitle',
        parent=styles['Heading1'],
        fontSize=16,
        alignment=TA_CENTER,
        spaceAfter=12,
    ))
    styles.add(ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontSize=13,
        spaceAfter=6,
        spaceBefore=12,
    ))
    styles.add(ParagraphStyle(
        'BoldLine',
        parent=styles['Normal'],
        fontSize=11,
        leading=14,
        spaceAfter=2,
    ))

    elements: list[Any] = []

    if rewritten_data:
        _build_structured_pdf(elements, rewritten_data, styles)
    elif raw_text:
        _build_plain_pdf(elements, raw_text, styles)
    else:
        elements.append(Paragraph('Данные для экспорта отсутствуют.', styles['Normal']))

    doc.build(elements)
    return buffer.getvalue()


def _build_structured_pdf(
    elements: list[Any],
    data: dict[str, Any],
    styles: Any,
) -> None:
    """Формируем PDF из структурированных JSON-данных LLM."""
    summary = data.get('summary', '')
    if summary:
        elements.append(Paragraph('Профессиональное резюме', styles['ResumeTitle']))
        elements.append(Paragraph(summary, styles['Normal']))
        elements.append(Spacer(1, 10))

    experience = data.get('experience', [])
    if experience:
        elements.append(Paragraph('Опыт работы', styles['SectionHeading']))
        for exp in experience:
            position = exp.get('position', '')
            company = exp.get('company', '')
            period = exp.get('period', '')
            if position:
                elements.append(Paragraph(f'<b>{position}</b>', styles['BoldLine']))
            parts = [p for p in (company, period) if p]
            if parts:
                elements.append(Paragraph(
                    f'<i>{" | ".join(parts)}</i>', styles['Normal'],
                ))
            for ach in exp.get('achievements', []):
                elements.append(Paragraph(f'• {ach}', styles['Normal']))
            elements.append(Spacer(1, 6))

    education = data.get('education', [])
    if education:
        elements.append(Paragraph('Образование', styles['SectionHeading']))
        for edu in education:
            institution = edu.get('institution', '')
            if institution:
                elements.append(Paragraph(f'<b>{institution}</b>', styles['BoldLine']))
            details = [
                d for d in (
                    edu.get('degree', ''),
                    edu.get('specialization', ''),
                    str(edu.get('year', '')),
                ) if d
            ]
            if details:
                elements.append(Paragraph(', '.join(details), styles['Normal']))

    skills = data.get('skills', [])
    if skills:
        elements.append(Paragraph('Навыки', styles['SectionHeading']))
        elements.append(Paragraph(', '.join(skills), styles['Normal']))


def _build_plain_pdf(elements: list[Any], text: str, styles: Any) -> None:
    """Формируем PDF из plain-text."""
    elements.append(Paragraph('Оптимизированное резюме', styles['ResumeTitle']))
    for paragraph in text.split('\n'):
        stripped = paragraph.strip()
        if stripped:
            elements.append(Paragraph(stripped, styles['Normal']))


# ---------------------------------------------------------------------------
# TXT generation
# ---------------------------------------------------------------------------

def generate_txt(rewritten_data: dict[str, Any] | None, *, raw_text: str | None = None) -> bytes:
    """Генерация plain-text версии резюме.

    Args:
        rewritten_data: Структурированные данные от LLM (JSON).
        raw_text: Fallback — текстовый ответ LLM.

    Returns:
        UTF-8 bytes.
    """
    if rewritten_data:
        return _build_structured_txt(rewritten_data).encode('utf-8')
    if raw_text:
        return raw_text.encode('utf-8')
    return 'Данные для экспорта отсутствуют.'.encode()


def _build_structured_txt(data: dict[str, Any]) -> str:
    """Plain-text из JSON-данных LLM."""
    lines: list[str] = []

    summary = data.get('summary', '')
    if summary:
        lines.append('ПРОФЕССИОНАЛЬНОЕ РЕЗЮМЕ')
        lines.append('=' * 40)
        lines.append(summary)
        lines.append('')

    experience = data.get('experience', [])
    if experience:
        lines.append('ОПЫТ РАБОТЫ')
        lines.append('-' * 40)
        for exp in experience:
            position = exp.get('position', '')
            company = exp.get('company', '')
            period = exp.get('period', '')
            if position:
                lines.append(position)
            parts = [p for p in (company, period) if p]
            if parts:
                lines.append(' | '.join(parts))
            for ach in exp.get('achievements', []):
                lines.append(f'  • {ach}')
            lines.append('')

    education = data.get('education', [])
    if education:
        lines.append('ОБРАЗОВАНИЕ')
        lines.append('-' * 40)
        for edu in education:
            institution = edu.get('institution', '')
            if institution:
                lines.append(institution)
            details = [
                d for d in (
                    edu.get('degree', ''),
                    edu.get('specialization', ''),
                    str(edu.get('year', '')),
                ) if d
            ]
            if details:
                lines.append(', '.join(details))
            lines.append('')

    skills = data.get('skills', [])
    if skills:
        lines.append('НАВЫКИ')
        lines.append('-' * 40)
        lines.append(', '.join(skills))

    return '\n'.join(lines)
