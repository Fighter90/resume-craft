"""ResumeCraft Streamlit Demo UI — для защиты ВКР (Phase 1).

Демонстрирует полный pipeline:
Upload (PDF/DOCX) → Parse → Match with Vacancy → AI Rewrite → Score → Export (DOCX)

Запуск:
    streamlit run streamlit_app/app.py
"""

from __future__ import annotations

import io
import json
import time
from pathlib import Path
from uuid import uuid4

import requests
import streamlit as st

# ── Конфигурация ────────────────────────────────────────────────────────────
API_BASE = 'http://localhost:8000/api/v1'

# ── Цветовая схема (из прототипов) ─────────────────────────────────────────
COLORS = {
    'primary': '#565ADD',
    'primary_hover': '#4338CA',
    'primary_light': '#EEF2FF',
    'success': '#10B981',
    'success_light': '#D1FAE5',
    'danger': '#EF4444',
    'danger_light': '#FEE2E2',
    'warning': '#F59E0B',
    'warning_light': '#FEF3C7',
    'bg_body': '#F3F4F6',
    'bg_card': '#FFFFFF',
    'text_main': '#111827',
    'text_secondary': '#6B7280',
}


def _inject_css() -> None:
    """Кастомные стили для приближения к прототипу."""
    st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');

    .stApp {
        font-family: 'Inter', sans-serif;
    }
    .main-header {
        background: linear-gradient(135deg, #4F46E5 0%, #7C3AED 100%);
        padding: 2rem;
        border-radius: 20px;
        color: white;
        margin-bottom: 1.5rem;
    }
    .main-header h1 { color: white; margin: 0; font-size: 2rem; }
    .main-header p { color: rgba(255,255,255,0.85); margin: 0.5rem 0 0 0; }
    .score-card {
        background: white;
        border-radius: 16px;
        padding: 1.5rem;
        box-shadow: 0 1px 3px rgba(0,0,0,0.1);
        text-align: center;
        height: 100%;
    }
    .score-card .value {
        font-size: 2.5rem;
        font-weight: 700;
        line-height: 1;
        margin: 0.5rem 0;
    }
    .score-card .label {
        font-size: 0.85rem;
        color: #6B7280;
        margin-bottom: 0.25rem;
    }
    .score-card .badge {
        display: inline-block;
        padding: 0.2rem 0.75rem;
        border-radius: 9999px;
        font-size: 0.75rem;
        font-weight: 600;
    }
    .badge-green { background: #D1FAE5; color: #059669; }
    .badge-yellow { background: #FEF3C7; color: #D97706; }
    .badge-red { background: #FEE2E2; color: #DC2626; }
    .badge-indigo { background: #EEF2FF; color: #4F46E5; }
    .keyword-pill {
        display: inline-block;
        padding: 0.3rem 0.75rem;
        border-radius: 9999px;
        font-size: 0.8rem;
        font-weight: 500;
        margin: 0.2rem;
        background: #DCFCE7;
        color: #166534;
    }
    .diff-original {
        background: #FEE2E2;
        padding: 0.2rem 0.4rem;
        border-radius: 4px;
    }
    .diff-improved {
        background: #DCFCE7;
        padding: 0.2rem 0.4rem;
        border-radius: 4px;
    }
    .step-active { color: #4F46E5; font-weight: 600; }
    .step-done { color: #10B981; }
    .step-pending { color: #9CA3AF; }
    .progress-section {
        background: white;
        border-radius: 16px;
        padding: 2rem;
        box-shadow: 0 1px 3px rgba(0,0,0,0.1);
    }
    </style>
    """, unsafe_allow_html=True)


def _get_score_color(score: float) -> str:
    """Цвет по значению оценки: ≥80 → зелёный, ≥60 → жёлтый, <60 → красный."""
    if score >= 80:
        return COLORS['success']
    if score >= 60:
        return COLORS['warning']
    return COLORS['danger']


def _get_badge_class(score: float) -> str:
    """CSS-класс бейджа по оценке."""
    if score >= 80:
        return 'badge-green'
    if score >= 60:
        return 'badge-yellow'
    return 'badge-red'


def _get_ats_color(rating: str) -> str:
    """Цвет ATS-рейтинга."""
    if rating in ('A+', 'A'):
        return COLORS['success']
    if rating in ('B+', 'B'):
        return COLORS['warning']
    return COLORS['danger']


# ── Состояние сессии ────────────────────────────────────────────────────────
def _init_session_state() -> None:
    """Инициализация session_state для wizard flow."""
    defaults = {
        'step': 'upload',       # upload → vacancy → model → processing → results
        'token': None,
        'user_email': None,
        'resume_id': None,
        'resume_text': None,
        'resume_filename': None,
        'vacancy_id': None,
        'vacancy_title': None,
        'vacancy_company': None,
        'vacancy_description': None,
        'model_name': 'gigachat-pro',
        'task_id': None,
        'result': None,
        'demo_mode': True,      # True = без бэкенда (демо-данные)
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v


# ── API-клиент ──────────────────────────────────────────────────────────────
def _api_headers() -> dict[str, str]:
    """JWT-заголовки."""
    headers: dict[str, str] = {}
    if st.session_state.token:
        headers['Authorization'] = f'Bearer {st.session_state.token}'
    return headers


def _api_request(
    method: str,
    path: str,
    *,
    json_data: dict | None = None,
    files: dict | None = None,
) -> dict | None:
    """Обёртка для API-запросов с обработкой ошибок."""
    url = f'{API_BASE}{path}'
    try:
        resp = requests.request(
            method, url,
            headers=_api_headers(),
            json=json_data,
            files=files,
            timeout=60,
        )
        if resp.status_code >= 400:
            error = resp.json() if resp.headers.get('content-type', '').startswith('application/json') else {}
            st.error(f'Ошибка API: {error.get("message", resp.status_code)}')
            return None
        if resp.status_code == 204:
            return {}
        return resp.json()
    except requests.ConnectionError:
        return None  # Demo mode will handle this


# ── Демо-данные ─────────────────────────────────────────────────────────────
DEMO_RESUME_TEXT = """Иванов Александр Сергеевич
Продакт-менеджер

Опыт работы:
Старший продакт-менеджер, ООО «ТехноСтарт», 2021–2024
- Управлял продуктом с MAU 200K
- Работал с командой разработки
- Запустил новый функционал

Продакт-менеджер, ООО «ДиджиталИнновация», 2019–2021
- Вёл backlog продукта
- Проводил интервью с пользователями
- Подготовил roadmap

Образование:
МГУ им. Ломоносова, Факультет ВМК, Прикладная математика, 2019

Навыки:
Python, SQL, Jira, Product Management, Аналитика"""

DEMO_OPTIMIZED = {
    'summary': (
        'Senior Product Manager с 5+ лет опыта в B2C и B2B SaaS. '
        'Руководил продуктами с MAU 500K+. Увеличил ключевые бизнес-метрики на 40% '
        'через Data-Driven подход и A/B тестирование. Опыт управления кросс-функциональными '
        'командами до 15 человек.'
    ),
    'experience': [
        {
            'position': 'Senior Product Manager',
            'company': 'ООО «ТехноСтарт»',
            'period': '2021–2024',
            'achievements': [
                'Руководил B2C-платформой с MAU 500K+, увеличил конверсию регистрации на 35% '
                'через серию A/B тестов (12 экспериментов за квартал)',
                'Внедрил систему OKR и Data-Driven принятие решений, что сократило Time-to-Market '
                'новых фич на 28% (с 6 до 4.3 недель)',
                'Управлял кросс-функциональной командой из 12 человек (Backend, Frontend, QA, Design), '
                'обеспечив 95% Sprint Delivery Rate',
            ],
        },
        {
            'position': 'Product Manager',
            'company': 'ООО «ДиджиталИнновация»',
            'period': '2019–2021',
            'achievements': [
                'Провёл 50+ Customer Development интервью, выявил 3 ключевых Pain Point, '
                'что привело к росту NPS с 32 до 58',
                'Разработал и реализовал Product Roadmap на 12 месяцев, '
                'приоритизируя по RICE-фреймворку',
                'Запустил MVP нового B2B-модуля за 8 недель, привлекшего 200+ корпоративных клиентов '
                'в первый квартал',
            ],
        },
    ],
    'education': [
        {
            'institution': 'МГУ им. М.В. Ломоносова',
            'degree': 'Специалист',
            'specialization': 'Прикладная математика и информатика',
            'year': 2019,
        },
    ],
    'skills': [
        'Product Strategy', 'A/B Testing', 'SQL', 'Python', 'Agile/Scrum',
        'Jira', 'Amplitude', 'Data-Driven', 'B2C SaaS', 'B2B SaaS',
        'OKR', 'RICE', 'Customer Development', 'Roadmap Planning',
        'Cross-functional Leadership', 'KPI', 'NPS', 'Unit Economics',
    ],
    'keywords_added': [
        'A/B Testing', 'Product Strategy', 'Data-Driven', 'Amplitude',
        'Cross-functional', 'KPI', 'B2C SaaS', 'Agile/Scrum',
        'MAU', 'Conversion', 'Sprint Delivery', 'B2B',
        'Roadmap', 'User Research', 'Metrics', 'OKR',
        'RICE', 'NPS',
    ],
}

DEMO_SCORES = {
    'match_score_before': 53.0,
    'match_score_after': 87.0,
    'ats_rating': 'A+',
    'keywords': {'weight': 40, 'score': 92},
    'experience': {'weight': 25, 'score': 85},
    'structure': {'weight': 20, 'score': 80},
    'readability': {'weight': 15, 'score': 88},
    'model_name': 'GigaChat Pro',
}

DEMO_VACANCY = {
    'title': 'Senior Product Manager',
    'company': 'Яндекс',
    'description': (
        'Мы ищем опытного Product Manager для работы над ключевыми B2C-продуктами. '
        'Вы будете отвечать за стратегию продукта, работу с данными и A/B тестирование. '
        'Требования: 5+ лет опыта в product management, опыт с B2C SaaS, '
        'знание SQL и аналитических инструментов (Amplitude, Mixpanel), '
        'опыт Agile/Scrum, понимание Unit Economics.'
    ),
    'key_skills': [
        'Product Strategy', 'A/B Testing', 'SQL', 'Agile/Scrum',
        'B2C SaaS', 'Amplitude', 'Data-Driven', 'Cross-functional',
    ],
    'city': 'Москва',
    'salary_from': 350000,
}


# ── Шаги Wizard ─────────────────────────────────────────────────────────────

def _render_sidebar() -> None:
    """Боковая панель навигации."""
    with st.sidebar:
        st.markdown("""
        <div style="display: flex; align-items: center; gap: 0.75rem; margin-bottom: 2rem;">
            <div style="width: 36px; height: 36px; background: linear-gradient(135deg, #4F46E5, #7C3AED);
                        border-radius: 10px; display: flex; align-items: center; justify-content: center;">
                <span style="color: white; font-weight: 700; font-size: 1.1rem;">R</span>
            </div>
            <span style="font-weight: 700; font-size: 1.25rem; color: #111827;">ResumeCraft</span>
        </div>
        """, unsafe_allow_html=True)

        st.caption('AI-реврайтер резюме')
        st.divider()

        steps = {
            'upload': ('📄', '1. Загрузка резюме'),
            'vacancy': ('🎯', '2. Выбор вакансии'),
            'model': ('🤖', '3. Выбор AI-модели'),
            'processing': ('⚡', '4. Оптимизация'),
            'results': ('📊', '5. Результаты'),
        }

        current = st.session_state.step
        step_order = list(steps.keys())
        current_idx = step_order.index(current)

        for idx, (key, (icon, label)) in enumerate(steps.items()):
            if idx < current_idx:
                st.markdown(f'✅ ~~{label}~~')
            elif idx == current_idx:
                st.markdown(f'**{icon} {label}**')
            else:
                st.markdown(f'⬜ {label}')

        st.divider()

        # Режим работы
        demo = st.toggle('Демо-режим', value=st.session_state.demo_mode)
        st.session_state.demo_mode = demo

        if demo:
            st.info('Работа без бэкенда — демо-данные', icon='ℹ️')
        else:
            st.caption('Подключение к API')
            if st.session_state.token:
                st.success(f'Авторизован: {st.session_state.user_email}', icon='✅')
            else:
                _render_login_form()


def _render_login_form() -> None:
    """Мини-форма логина в sidebar."""
    with st.form('login_form'):
        email = st.text_input('Email', placeholder='user@example.com')
        password = st.text_input('Пароль', type='password')
        if st.form_submit_button('Войти', use_container_width=True):
            data = _api_request('POST', '/auth/login', json_data={
                'email': email,
                'password': password,
            })
            if data and 'access_token' in data:
                st.session_state.token = data['access_token']
                st.session_state.user_email = email
                st.rerun()
            else:
                st.error('Ошибка авторизации')


def _render_upload() -> None:
    """Шаг 1: Загрузка резюме."""
    st.markdown("""
    <div class="main-header">
        <h1>📄 Загрузите ваше резюме</h1>
        <p>Мы поддерживаем форматы PDF и DOCX до 10 МБ</p>
    </div>
    """, unsafe_allow_html=True)

    col1, col2 = st.columns([2, 1])

    with col1:
        uploaded = st.file_uploader(
            'Выберите файл резюме',
            type=['pdf', 'docx'],
            help='Поддерживаемые форматы: PDF, DOCX. Максимальный размер: 10 МБ.',
        )

        if uploaded:
            st.session_state.resume_filename = uploaded.name
            file_size_mb = len(uploaded.getvalue()) / (1024 * 1024)
            st.success(f'Файл загружен: **{uploaded.name}** ({file_size_mb:.1f} МБ)', icon='✅')

            if st.session_state.demo_mode:
                st.session_state.resume_text = DEMO_RESUME_TEXT
                st.session_state.resume_id = str(uuid4())
            else:
                # Отправить на бэкенд
                data = _api_request('POST', '/resumes/upload', files={
                    'file': (uploaded.name, uploaded.getvalue(), uploaded.type),
                })
                if data:
                    st.session_state.resume_id = data['id']
                    st.session_state.resume_text = data.get('raw_text', '')

            if st.session_state.resume_id:
                if st.button('Далее → Выбор вакансии', type='primary', use_container_width=True):
                    st.session_state.step = 'vacancy'
                    st.rerun()

        # Демо-кнопка
        st.divider()
        if st.button('🎬 Загрузить демо-резюме', use_container_width=True):
            st.session_state.resume_text = DEMO_RESUME_TEXT
            st.session_state.resume_id = str(uuid4())
            st.session_state.resume_filename = 'demo_resume.pdf'
            st.session_state.step = 'vacancy'
            st.rerun()

    with col2:
        st.markdown('### 💡 Советы')
        st.markdown("""
        - Используйте **актуальное** резюме
        - PDF с текстовым слоем даёт лучший результат
        - Не загружайте сканы — качество OCR ниже
        - Убедитесь, что файл < 10 МБ
        """)


def _render_vacancy() -> None:
    """Шаг 2: Выбор вакансии."""
    st.markdown("""
    <div class="main-header">
        <h1>🎯 Выберите целевую вакансию</h1>
        <p>Укажите вакансию, под которую нужно оптимизировать резюме</p>
    </div>
    """, unsafe_allow_html=True)

    tab1, tab2, tab3 = st.tabs(['🔍 Поиск hh.ru', '🔗 Вставить URL', '✏️ Ввести вручную'])

    with tab1:
        if st.session_state.demo_mode:
            st.info('В демо-режиме поиск hh.ru недоступен. Используйте демо-вакансию ниже.')
            if st.button('📋 Использовать демо-вакансию (Senior PM @ Яндекс)', use_container_width=True):
                st.session_state.vacancy_title = DEMO_VACANCY['title']
                st.session_state.vacancy_company = DEMO_VACANCY['company']
                st.session_state.vacancy_description = DEMO_VACANCY['description']
                st.session_state.vacancy_id = str(uuid4())
                st.session_state.step = 'model'
                st.rerun()
        else:
            col1, col2 = st.columns(2)
            with col1:
                search_text = st.text_input('Должность', placeholder='Product Manager')
            with col2:
                city = st.selectbox('Город', ['Москва', 'Санкт-Петербург', 'Удалённо', 'Другой'])

            if st.button('Найти', type='primary'):
                data = _api_request('GET', '/vacancies/search', json_data={
                    'text': search_text,
                })
                if data and data.get('items'):
                    for item in data['items'][:6]:
                        with st.container():
                            st.markdown(f"**{item['title']}** @ {item.get('company', 'N/A')}")
                            if st.button('Выбрать', key=f"vac_{item.get('id', '')}"):
                                st.session_state.vacancy_id = item['id']
                                st.session_state.vacancy_title = item['title']
                                st.session_state.step = 'model'
                                st.rerun()

    with tab2:
        url = st.text_input('Ссылка на вакансию hh.ru', placeholder='https://hh.ru/vacancy/12345678')
        if url and st.button('Загрузить вакансию', type='primary'):
            if st.session_state.demo_mode:
                st.session_state.vacancy_title = DEMO_VACANCY['title']
                st.session_state.vacancy_company = DEMO_VACANCY['company']
                st.session_state.vacancy_description = DEMO_VACANCY['description']
                st.session_state.vacancy_id = str(uuid4())
                st.session_state.step = 'model'
                st.rerun()
            else:
                data = _api_request('POST', '/vacancies/from-url', json_data={'url': url})
                if data:
                    st.session_state.vacancy_id = data['id']
                    st.session_state.vacancy_title = data['title']
                    st.session_state.step = 'model'
                    st.rerun()

    with tab3:
        with st.form('manual_vacancy'):
            title = st.text_input('Название должности *', placeholder='Senior Product Manager')
            company = st.text_input('Компания', placeholder='Яндекс')
            description = st.text_area(
                'Описание вакансии и требования *',
                height=200,
                placeholder='Опишите требования к кандидату, ключевые навыки, обязанности...',
            )
            skills = st.text_input(
                'Ключевые навыки (через запятую)',
                placeholder='Product Strategy, A/B Testing, SQL, Agile',
            )

            if st.form_submit_button('Далее → Выбор модели', type='primary', use_container_width=True):
                if not title or not description:
                    st.error('Заполните обязательные поля')
                else:
                    if st.session_state.demo_mode:
                        st.session_state.vacancy_id = str(uuid4())
                    else:
                        data = _api_request('POST', '/vacancies/manual', json_data={
                            'title': title,
                            'company': company,
                            'description': description,
                            'key_skills': [s.strip() for s in skills.split(',') if s.strip()] if skills else [],
                        })
                        if data:
                            st.session_state.vacancy_id = data['id']

                    st.session_state.vacancy_title = title
                    st.session_state.vacancy_company = company
                    st.session_state.vacancy_description = description
                    st.session_state.step = 'model'
                    st.rerun()

    # Кнопка назад
    st.divider()
    if st.button('← Назад к загрузке'):
        st.session_state.step = 'upload'
        st.rerun()


def _render_model_selection() -> None:
    """Шаг 3: Выбор AI-модели."""
    st.markdown("""
    <div class="main-header">
        <h1>🤖 Выберите AI-модель</h1>
        <p>Различные модели подходят для разных задач</p>
    </div>
    """, unsafe_allow_html=True)

    models = [
        {
            'name': 'gigachat-pro',
            'display': 'GigaChat Pro',
            'provider': 'Сбер',
            'description': '#1 MERA для русского языка. Лучшее качество для русскоязычных резюме.',
            'speed': '10–20 сек',
            'quality': '⭐⭐⭐⭐⭐',
            'icon': '🇷🇺',
            'recommended': True,
        },
        {
            'name': 'groq',
            'display': 'Llama 3.3 70B (Groq)',
            'provider': 'Meta / Groq',
            'description': 'Быстрая бесплатная модель. Хорошо работает с IT-резюме.',
            'speed': '5–10 сек',
            'quality': '⭐⭐⭐⭐',
            'icon': '🦙',
            'recommended': False,
        },
        {
            'name': 'openai',
            'display': 'GPT-4o mini',
            'provider': 'OpenAI',
            'description': 'Резервная модель. Хорошее качество для англоязычного контента.',
            'speed': '10–15 сек',
            'quality': '⭐⭐⭐⭐',
            'icon': '🧠',
            'recommended': False,
        },
    ]

    cols = st.columns(len(models))
    for col, model in zip(cols, models):
        with col:
            badge = ' 🏆 Рекомендуем' if model['recommended'] else ''
            st.markdown(f"""
            <div class="score-card" style="text-align: left;">
                <div style="font-size: 2rem;">{model['icon']}</div>
                <div style="font-weight: 700; font-size: 1.1rem; margin: 0.5rem 0;">
                    {model['display']}{badge}
                </div>
                <div style="color: #6B7280; font-size: 0.85rem; margin-bottom: 0.5rem;">
                    {model['provider']}
                </div>
                <div style="font-size: 0.9rem; margin-bottom: 0.75rem;">
                    {model['description']}
                </div>
                <div style="font-size: 0.85rem;">
                    ⚡ {model['speed']} &nbsp;|&nbsp; {model['quality']}
                </div>
            </div>
            """, unsafe_allow_html=True)

            if st.button(
                'Выбрать' if not model['recommended'] else '✓ Выбрать (рекомендуем)',
                key=f"model_{model['name']}",
                type='primary' if model['recommended'] else 'secondary',
                use_container_width=True,
            ):
                st.session_state.model_name = model['name']
                st.session_state.step = 'processing'
                st.rerun()

    st.divider()
    if st.button('← Назад к вакансии'):
        st.session_state.step = 'vacancy'
        st.rerun()


def _render_processing() -> None:
    """Шаг 4: Обработка (анимация прогресса)."""
    st.markdown("""
    <div class="main-header" style="text-align: center;">
        <h1>⚡ Оптимизируем ваше резюме...</h1>
        <p>{model} анализирует и адаптирует контент</p>
    </div>
    """.format(model=st.session_state.model_name.replace('-', ' ').title()), unsafe_allow_html=True)

    steps = [
        ('📄', 'Загрузка документа', 'Парсинг и извлечение текста'),
        ('🔍', 'Анализ вакансии', 'Извлечение ключевых требований'),
        ('🤖', 'AI оптимизация', f'{st.session_state.model_name} адаптирует контент'),
        ('✅', 'Финализация', 'Проверка ATS и форматирование'),
    ]

    progress_bar = st.progress(0)
    status_area = st.empty()
    step_list = st.empty()

    if st.session_state.demo_mode:
        # Демо-анимация
        for i, (icon, title, desc) in enumerate(steps):
            pct = int((i + 1) / len(steps) * 100)
            progress_bar.progress(pct)

            steps_html = ''
            for j, (s_icon, s_title, s_desc) in enumerate(steps):
                if j < i:
                    steps_html += f'<div class="step-done">✅ {s_title} — {s_desc}</div>'
                elif j == i:
                    steps_html += f'<div class="step-active">⏳ {s_title} — {s_desc}</div>'
                else:
                    steps_html += f'<div class="step-pending">⬜ {s_title} — {s_desc}</div>'

            step_list.markdown(f"""
            <div class="progress-section">
                {steps_html}
            </div>
            """, unsafe_allow_html=True)

            time.sleep(1.5)

        # Готово!
        progress_bar.progress(100)
        st.session_state.result = {
            **DEMO_SCORES,
            'optimized': DEMO_OPTIMIZED,
        }
        st.balloons()
        time.sleep(0.5)
        st.session_state.step = 'results'
        st.rerun()
    else:
        # Реальный API-запрос
        data = _api_request('POST', '/rewrite', json_data={
            'resume_id': st.session_state.resume_id,
            'vacancy_id': st.session_state.vacancy_id,
            'model_name': st.session_state.model_name,
        })
        if data:
            task_id = data['task_id']
            while True:
                status = _api_request('GET', f'/rewrite/{task_id}/status')
                if not status:
                    break
                progress_bar.progress(status.get('progress', 0))
                status_area.info(f"Шаг: {status.get('step', '...')}")

                if status.get('status') in ('completed', 'failed'):
                    break
                time.sleep(2)

            result = _api_request('GET', f'/rewrite/{task_id}/result')
            if result:
                st.session_state.result = result
                st.session_state.step = 'results'
                st.rerun()


def _render_results() -> None:
    """Шаг 5: Результаты оптимизации."""
    result = st.session_state.result
    if not result:
        st.error('Результаты не найдены')
        return

    vacancy_label = st.session_state.vacancy_title or 'Вакансия'
    company_label = st.session_state.vacancy_company or ''
    company_str = f' @ {company_label}' if company_label else ''

    st.markdown(f"""
    <div class="main-header">
        <h1>📊 Результаты оптимизации</h1>
        <p>{vacancy_label}{company_str}</p>
    </div>
    """, unsafe_allow_html=True)

    # ── Score карточки ──
    score_before = result.get('match_score_before', 53)
    score_after = result.get('match_score_after', 87)
    improvement = score_after - score_before
    ats = result.get('ats_rating', 'A+')
    keywords_added = result.get('optimized', {}).get('keywords_added', [])
    model = result.get('model_name', 'GigaChat Pro')

    c1, c2, c3, c4 = st.columns(4)

    with c1:
        color = _get_score_color(score_after)
        badge_cls = _get_badge_class(score_after)
        st.markdown(f"""
        <div class="score-card">
            <div class="label">Match Score</div>
            <div class="value" style="color: {color};">{score_after:.0f}%</div>
            <div class="badge {badge_cls}">+{improvement:.0f}% улучшение</div>
        </div>
        """, unsafe_allow_html=True)

    with c2:
        ats_color = _get_ats_color(ats)
        st.markdown(f"""
        <div class="score-card">
            <div class="label">ATS совместимость</div>
            <div class="value" style="color: {ats_color};">{ats}</div>
            <div class="badge badge-green">Все секции оптимальны</div>
        </div>
        """, unsafe_allow_html=True)

    with c3:
        kw_count = len(keywords_added)
        st.markdown(f"""
        <div class="score-card">
            <div class="label">Ключевые слова</div>
            <div class="value" style="color: #565ADD;">{kw_count}</div>
            <div class="badge badge-green">+{kw_count} добавлено</div>
        </div>
        """, unsafe_allow_html=True)

    with c4:
        st.markdown(f"""
        <div class="score-card">
            <div class="label">AI модель</div>
            <div class="value" style="font-size: 1.5rem; color: #111827;">{model}</div>
            <div class="badge badge-indigo">⚡ Основная</div>
        </div>
        """, unsafe_allow_html=True)

    st.markdown('<br>', unsafe_allow_html=True)

    # ── Разбивка Match Score ──
    st.subheader('Разбивка Match Score')
    components = [
        ('Ключевые слова', result.get('keywords', {}).get('weight', 40),
         result.get('keywords', {}).get('score', 92)),
        ('Опыт', result.get('experience', {}).get('weight', 25),
         result.get('experience', {}).get('score', 85)),
        ('Структура', result.get('structure', {}).get('weight', 20),
         result.get('structure', {}).get('score', 80)),
        ('Читаемость', result.get('readability', {}).get('weight', 15),
         result.get('readability', {}).get('score', 88)),
    ]

    for label, weight, score in components:
        col_label, col_bar, col_val = st.columns([2, 5, 1])
        with col_label:
            st.markdown(f'**{label}** ({weight}%)')
        with col_bar:
            st.progress(score / 100)
        with col_val:
            color = _get_score_color(score)
            st.markdown(f'<span style="color: {color}; font-weight: 600;">{score}%</span>',
                        unsafe_allow_html=True)

    st.markdown('<br>', unsafe_allow_html=True)

    # ── Diff сравнение ──
    st.subheader('Сравнение: Оригинал → Оптимизировано')

    col_orig, col_opt = st.columns(2)

    with col_orig:
        st.markdown(f"""
        <div style="background: white; border-radius: 12px; padding: 1.5rem;
                    box-shadow: 0 1px 3px rgba(0,0,0,0.1);">
            <div style="display: flex; justify-content: space-between; align-items: center;
                        margin-bottom: 1rem;">
                <strong>Оригинал</strong>
                <span class="badge badge-red">{score_before:.0f}%</span>
            </div>
            <div style="font-size: 0.9rem; line-height: 1.6; color: #374151;">
                {st.session_state.resume_text or 'Текст оригинала'}
            </div>
        </div>
        """, unsafe_allow_html=True)

    with col_opt:
        optimized = result.get('optimized', {})
        if isinstance(optimized, dict):
            summary = optimized.get('summary', '')
            experience = optimized.get('experience', [])
            skills = optimized.get('skills', [])

            exp_html = ''
            for exp in experience[:2]:
                achievements = ''.join(
                    f'<li>{a}</li>' for a in exp.get('achievements', [])
                )
                exp_html += f"""
                <div style="margin-bottom: 1rem;">
                    <strong>{exp.get('position', '')}</strong> — {exp.get('company', '')}
                    <span style="color: #6B7280;">({exp.get('period', '')})</span>
                    <ul style="margin: 0.25rem 0; padding-left: 1.25rem;">{achievements}</ul>
                </div>
                """

            skills_html = ', '.join(f'<span class="diff-improved">{s}</span>' for s in skills[:10])

            st.markdown(f"""
            <div style="background: white; border-radius: 12px; padding: 1.5rem;
                        border-left: 3px solid #10B981;
                        box-shadow: 0 1px 3px rgba(0,0,0,0.1);">
                <div style="display: flex; justify-content: space-between; align-items: center;
                            margin-bottom: 1rem;">
                    <strong>Оптимизировано</strong>
                    <span class="badge badge-green">{score_after:.0f}%</span>
                </div>
                <div style="font-size: 0.9rem; line-height: 1.6; color: #374151;">
                    <p><span class="diff-improved">{summary}</span></p>
                    {exp_html}
                    <p><strong>Навыки:</strong> {skills_html}</p>
                </div>
            </div>
            """, unsafe_allow_html=True)

    st.markdown('<br>', unsafe_allow_html=True)

    # ── Добавленные ключевые слова ──
    if keywords_added:
        st.subheader('Добавленные ключевые слова')
        kw_html = ''.join(f'<span class="keyword-pill">{kw}</span>' for kw in keywords_added)
        st.markdown(f'<div>{kw_html}</div>', unsafe_allow_html=True)

    st.markdown('<br>', unsafe_allow_html=True)

    # ── Действия ──
    col_a, col_b, col_c = st.columns(3)

    with col_a:
        if st.button('✏️ Редактировать', use_container_width=True):
            st.info('Редактор будет доступен в Phase 2 (React SPA)')

    with col_b:
        # Генерация DOCX для скачивания
        if st.button('📥 Экспорт DOCX', type='primary', use_container_width=True):
            if st.session_state.demo_mode:
                _export_demo_docx(result)
            else:
                task_id = st.session_state.task_id or st.session_state.result.get('id', '')
                resp = requests.get(
                    f'{API_BASE}/export/{task_id}/docx',
                    headers=_api_headers(),
                    timeout=30,
                )
                if resp.status_code == 200:
                    st.download_button(
                        '💾 Скачать DOCX', data=resp.content,
                        file_name='resume_optimized.docx',
                        mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                    )

    with col_c:
        if st.button('🔄 Новая оптимизация', use_container_width=True):
            for key in ['resume_id', 'resume_text', 'vacancy_id', 'vacancy_title',
                        'vacancy_company', 'vacancy_description', 'task_id', 'result']:
                st.session_state[key] = None
            st.session_state.step = 'upload'
            st.rerun()


def _export_demo_docx(result: dict) -> None:
    """Генерация DOCX из демо-данных (без бэкенда)."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt

        doc = Document()
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)

        optimized = result.get('optimized', {})

        # Заголовок
        doc.add_heading('Иванов Александр Сергеевич', level=1)
        doc.add_paragraph(optimized.get('summary', ''))

        # Опыт
        doc.add_heading('Опыт работы', level=2)
        for exp in optimized.get('experience', []):
            doc.add_heading(
                f"{exp.get('position', '')} — {exp.get('company', '')} ({exp.get('period', '')})",
                level=3,
            )
            for ach in exp.get('achievements', []):
                doc.add_paragraph(f'• {ach}')

        # Образование
        doc.add_heading('Образование', level=2)
        for edu in optimized.get('education', []):
            doc.add_paragraph(
                f"{edu.get('institution', '')}, {edu.get('specialization', '')} ({edu.get('year', '')})"
            )

        # Навыки
        doc.add_heading('Навыки', level=2)
        skills = optimized.get('skills', [])
        doc.add_paragraph(', '.join(skills))

        # Сохранение в буфер
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        st.download_button(
            '💾 Скачать DOCX',
            data=buf.getvalue(),
            file_name='resume_optimized.docx',
            mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )

    except ImportError:
        st.warning('python-docx не установлен. Экспорт недоступен.')


# ── Main ────────────────────────────────────────────────────────────────────

def main() -> None:
    """Точка входа Streamlit-приложения."""
    st.set_page_config(
        page_title='ResumeCraft — AI-реврайтер резюме',
        page_icon='📄',
        layout='wide',
        initial_sidebar_state='expanded',
    )

    _init_session_state()
    _inject_css()
    _render_sidebar()

    # Wizard routing
    step = st.session_state.step
    if step == 'upload':
        _render_upload()
    elif step == 'vacancy':
        _render_vacancy()
    elif step == 'model':
        _render_model_selection()
    elif step == 'processing':
        _render_processing()
    elif step == 'results':
        _render_results()


if __name__ == '__main__':
    main()
